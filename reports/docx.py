"""Word (.docx) report generation.

NOTE on tooling: the Day-1 contract named docxtpl (Jinja-in-a-.docx-template).
That needs an actual .docx template file authored in Word, which isn't
practical to hand-build outside Word/LibreOffice. This module produces the
identical content and layout to the PDF using python-docx directly instead
-- still a fully editable Word document, just built programmatically rather
than from a template file. If you'd rather use docxtpl: open report.html's
structure as a guide, build report_template.docx in Word with {{ }} tags
matching build_view()'s keys, drop it in reports/templates/, and swap this
file's body for a docxtpl DocxTemplate(...).render(view) call.

pip install python-docx
"""
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt, RGBColor

from .common import build_view

NAVY = RGBColor(0x0B, 0x3D, 0x91)
GREY = RGBColor(0x66, 0x66, 0x66)


def _colour(hexstr: str) -> RGBColor:
    hexstr = hexstr.lstrip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


def _kv_table(doc: Document, rows: list[tuple[str, str]]):
    t = doc.add_table(rows=0, cols=2)
    t.autofit = True
    for k, v in rows:
        row = t.add_row().cells
        row[0].text = str(k)
        row[0].paragraphs[0].runs[0].font.color.rgb = GREY
        row[0].paragraphs[0].runs[0].font.size = Pt(9)
        row[1].text = "-" if v is None else str(v)
        row[1].paragraphs[0].runs[0].font.size = Pt(9)
    return t


def _heading(doc: Document, text: str):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = NAVY
        run.font.size = Pt(12)


def generate_docx(context: dict) -> bytes:
    """context = backend.services.reporting.get_report_context(db, session, lang)."""
    v = build_view(context)
    L = v["L"]
    doc = Document()

    # -- header --------------------------------------------------------
    hdr = doc.add_paragraph()
    hdr.add_run(v["lab"]["name"] if v["lab"] else "ScaleSaathi Testing Laboratory").bold = True
    if v["lab"] and v["lab"].get("address"):
        doc.add_paragraph(v["lab"]["address"]).runs[0].font.size = Pt(9)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta.add_run(f"{L['report_no']}: {v['report'].get('report_no') or '-'}    "
                 f"{L['date']}: {v['report'].get('approved_at') or '-'}").font.size = Pt(9)

    title = doc.add_heading(L["doc_title"], level=1)
    title.runs[0].font.color.rgb = NAVY

    # -- instrument + conditions ----------------------------------------
    _heading(doc, L["instrument_details"])
    _kv_table(doc, [
        (L["manufacturer"], v["manufacturer"]["name"]),
        (L["model"], v["instrument"]["model"]),
        (L["serial_no"], v["instrument"]["serial_no"]),
        (L["accuracy_class"], v["instrument"]["accuracy_class"]),
        (L["max_capacity"], f"{v['instrument']['max_capacity']} g"),
        (L["min_capacity"], f"{v['instrument']['min_capacity']} g"),
        (L["scale_interval"], f"{v['instrument']['e']} g"),
        (L["verification_interval"], f"{v['instrument']['d']} g"),
        (L["max_tare"], f"{v['instrument']['max_tare']} g"),
    ])

    _heading(doc, L["test_conditions"])
    cond = v["session"]["conditions"]
    _kv_table(doc, [
        (L["location"], v["session"].get("location")),
        (L["tested_on"], v["session"].get("tested_at")),
        (L["tested_by"], v["tester"]["full_name"] if v["tester"] else "-"),
        (L["temperature"], f"{cond.get('temperature_c')} °C"),
        (L["humidity"], f"{cond.get('humidity_pct')} %RH"),
        (L["pressure"], f"{cond.get('pressure_hpa')} hPa"),
    ])

    # -- callouts --------------------------------------------------------
    if v["trap_rows"]:
        p = doc.add_paragraph()
        r = p.add_run(f"\u26a0 {L['rounding_trap_title']}. {L['rounding_trap_body']}")
        r.bold = True
        r.font.color.rgb = _colour("#b42318")
    if v["marginal_rows"]:
        p = doc.add_paragraph()
        r = p.add_run(f"{L['marginal_title']}. {L['marginal_body']}")
        r.bold = True
        r.font.color.rgb = _colour("#b54708")

    # -- results table -----------------------------------------------------
    _heading(doc, L["results"])
    cols = [L["test"], L["load"], L["error"], L["mpe"], L["utilisation"],
            L["naive_method"], L["result"], L["clause"], L["explanation"]]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for r in v["results"]:
        cells = t.add_row().cells
        util = "-" if r["utilisation"] is None else f"{r['utilisation'] * 100:.0f}%"
        naive = "-" if not r["naive_result"] else f"{r['naive_error']} \u2192 {r['naive_result']}"
        values = [r["test"] + (" \u26a0" if r["marginal"] else ""),
                  r["load"], r["error"], r["mpe"], util, naive, r["result"],
                  r["clause"] or "-", r["explanation"] or "-"]
        for i, val in enumerate(values):
            cells[i].text = "-" if val is None else str(val)
            if i == 6:  # result column
                cells[i].paragraphs[0].runs[0].font.color.rgb = _colour(r["colour"])
                cells[i].paragraphs[0].runs[0].bold = True
        if r["is_rounding_trap"]:
            for c in cells:
                c.paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW

    # -- overall verdict banner --------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{L['overall_verdict']}: {v['overall']}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = _colour(v["overall_colour"])

    # -- attachments ---------------------------------------------------------
    if v["attachments"]:
        _heading(doc, L["attachments"])
        for a in v["attachments"]:
            if a.get("data"):
                try:
                    doc.add_picture(io.BytesIO(a["data"]), width=Cm(6))
                    cap = doc.add_paragraph(a.get("caption") or a.get("filename"))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.runs[0].font.size = Pt(8)
                except Exception:
                    doc.add_paragraph(f"[image: {a.get('filename')}]")

    if v["annexes"]:
        _heading(doc, L["annexes"])
        for a in v["annexes"]:
            doc.add_paragraph(f"{a['filename']} ({a['kind']})", style="List Bullet")

    # -- signature + QR -------------------------------------------------------
    doc.add_paragraph()
    sig_table = doc.add_table(rows=1, cols=2)
    left, right = sig_table.rows[0].cells
    ab = v["approved_by"]
    left.paragraphs[0].add_run(
        f"{ab['full_name'] if ab else '-'}\n"
        f"{ab.get('designation') or '' if ab else ''}\n"
        f"{L['approved_on']}: {v['report'].get('approved_at') or '-'}"
    ).font.size = Pt(9)
    if v.get("qr_data_uri"):
        import base64
        qr_bytes = base64.b64decode(v["qr_data_uri"].split(",", 1)[1])
        right.paragraphs[0].add_run().add_picture(io.BytesIO(qr_bytes), width=Cm(2.5))

    foot = doc.add_paragraph()
    r = foot.add_run(f"{L['verify_note']} {v.get('verify_url') or ''}\n{L['footer']}")
    r.font.size = Pt(7.5)
    r.font.color.rgb = GREY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
